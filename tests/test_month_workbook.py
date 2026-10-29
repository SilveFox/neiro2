# -*- coding: utf-8 -*-
"""Tests for empty month workbook creation."""
from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from core import config
from core.workbook_factory import (
    WorkbookExistsError,
    activate_workbook_name,
    create_month_workbook,
    list_month_workbooks,
    sanitize_workbook,
    sanitize_workbook_title,
    written_log_for,
)


class MonthWorkbookTests(unittest.TestCase):
    def test_create_empty_month(self):
        self.assertTrue(config.SOURCE_WORKBOOK.exists())
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            old_dir = config.WORKBOOKS_DIR
            old_meta = config.ACTIVE_WORKBOOK_META
            config.WORKBOOKS_DIR = tmp_path / "workbooks"
            config.WORKBOOKS_DIR.mkdir()
            config.ACTIVE_WORKBOOK_META = tmp_path / "active_workbook.json"
            try:
                path = create_month_workbook(2026, 7, activate=True)
                self.assertTrue(path.exists())
                self.assertEqual(path.name, "2026-07.xlsx")
                self.assertTrue(written_log_for(path).exists())

                wb = openpyxl.load_workbook(path, data_only=False)
                try:
                    self.assertIn("Данные", wb.sheetnames)
                    self.assertIn("01", wb.sheetnames)
                    self.assertIn("Шаблон", wb.sheetnames)
                    ws01 = wb["01"]
                    self.assertEqual(ws01.cell(1, 2).value, "01.07.2026")
                    # Work tables cleared; employee list kept on the right
                    self.assertIsNone(ws01.cell(3, 4).value)  # D header gone
                    self.assertIsNone(ws01.cell(4, 4).value)
                    # No leftover borders from template tables
                    b = ws01.cell(4, 4).border

                    def _no_side(side):
                        return side is None or side.style is None or side.style == "none"

                    self.assertTrue(
                        _no_side(b.left)
                        and _no_side(b.right)
                        and _no_side(b.top)
                        and _no_side(b.bottom)
                    )
                    self.assertEqual(ws01.cell(1, 11).value, "ФИО")
                    self.assertIsNotNone(ws01.cell(3, 11).value)
                    qty_values = [
                        ws01.cell(r, config.COL_QTY).value
                        for r in range(1, min(200, (ws01.max_row or 1) + 1))
                    ]
                    numeric_qty = [v for v in qty_values if isinstance(v, (int, float))]
                    self.assertEqual(numeric_qty, [])

                    calc_ws = wb["Расчет"]
                    self.assertEqual(calc_ws.cell(2, 6).value.date(), date(2026, 7, 1))
                    self.assertEqual(calc_ws.cell(2, 36).value.date(), date(2026, 7, 31))
                finally:
                    wb.close()

                with self.assertRaises(WorkbookExistsError):
                    create_month_workbook(2026, 7, activate=False)

                meta = json.loads(config.ACTIVE_WORKBOOK_META.read_text(encoding="utf-8"))
                self.assertEqual(meta["year_month"], "2026-07")
                self.assertTrue(Path(meta["path"]).exists())

                # Template external links must be stripped (they break Excel after rewrite)
                wb2 = openpyxl.load_workbook(path)
                try:
                    self.assertFalse(getattr(wb2, "_external_links", None))
                    for dn in wb2.defined_names.values():
                        text = getattr(dn, "attr_text", "") or ""
                        self.assertNotRegex(text, r"\[\d+\]")
                finally:
                    wb2.close()
            finally:
                config.WORKBOOKS_DIR = old_dir
                config.ACTIVE_WORKBOOK_META = old_meta

    def test_ensure_creates_missing_month_and_write_uses_report_date(self):
        """Write goes to YYYY-MM of report_date, not the previously active book."""
        self.assertTrue(config.SOURCE_WORKBOOK.exists())
        from core.excel_writer import append_calc_to_workbook
        from core.pricing import calculate_report
        from core.workbook_factory import (
            ensure_month_workbook,
            get_active_meta,
            month_workbook_path,
        )

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            old_dir = config.WORKBOOKS_DIR
            old_meta = config.ACTIVE_WORKBOOK_META
            config.WORKBOOKS_DIR = tmp_path / "workbooks"
            config.WORKBOOKS_DIR.mkdir()
            config.ACTIVE_WORKBOOK_META = tmp_path / "active_workbook.json"
            try:
                july, created_july = ensure_month_workbook(2026, 7, activate=True)
                self.assertTrue(created_july)
                self.assertTrue(july.exists())
                self.assertEqual(get_active_meta()["year_month"], "2026-07")

                report = {
                    "report_date": "2026-08-10",
                    "source_text_hash": "month-by-date-hash",
                    "workers": [
                        {
                            "raw_worker_name": "Карандак В.Е.",
                            "members": ["Карандак В.Е."],
                            "performed_jobs": [
                                {
                                    "raw_task_name": "обжим рж45",
                                    "job_description": "обжим рж45",
                                    "service_type": "Интернет",
                                    "volume": 1,
                                    "unit": "шт",
                                }
                            ],
                        }
                    ],
                }
                calc = calculate_report(report)
                result = append_calc_to_workbook(calc, skip_duplicates=True)
                self.assertTrue(result["ok"])
                self.assertTrue(result["workbook_created"])
                self.assertEqual(result["year_month"], "2026-08")
                self.assertEqual(result["day_sheet"], "10")
                aug = month_workbook_path(2026, 8)
                self.assertTrue(aug.exists())
                self.assertEqual(get_active_meta()["year_month"], "2026-08")
                self.assertNotEqual(Path(result["workbook"]).resolve(), july.resolve())
            finally:
                config.WORKBOOKS_DIR = old_dir
                config.ACTIVE_WORKBOOK_META = old_meta

    def test_sanitize_drops_external_names(self):
        self.assertTrue(config.SOURCE_WORKBOOK.exists())
        wb = openpyxl.load_workbook(config.SOURCE_WORKBOOK)
        try:
            before = list(getattr(wb, "_external_links", []) or [])
            self.assertTrue(before)  # template has external links
            sanitize_workbook(wb)
            self.assertEqual(list(getattr(wb, "_external_links", []) or []), [])
            for dn in wb.defined_names.values():
                text = getattr(dn, "attr_text", "") or ""
                self.assertNotRegex(text, r"\[\d+\]")
        finally:
            wb.close()

    def test_create_with_custom_title(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            old_dir = config.WORKBOOKS_DIR
            old_meta = config.ACTIVE_WORKBOOK_META
            old_src = config.SOURCE_WORKBOOK
            config.WORKBOOKS_DIR = tmp_path / "workbooks"
            config.WORKBOOKS_DIR.mkdir()
            config.ACTIVE_WORKBOOK_META = tmp_path / "active_workbook.json"
            # Minimal template so test does not depend on docs/*.xlsx presence
            src = tmp_path / "template.xlsx"
            wb = openpyxl.Workbook()
            wb.active.title = "Данные"
            wb.create_sheet("Расчет")
            wb.create_sheet("Шаблон")
            for day in range(1, 32):
                wb.create_sheet(f"{day:02d}")
            wb.save(src)
            wb.close()
            config.SOURCE_WORKBOOK = src
            try:
                self.assertEqual(sanitize_workbook_title("  Гарантия / север "), "Гарантия_север")
                path = create_month_workbook(2026, 11, activate=True, title="Гарантия")
                self.assertEqual(path.name, "Гарантия_2026-11.xlsx")
                items = list_month_workbooks()
                self.assertEqual(len(items), 1)
                self.assertEqual(items[0]["title"], "Гарантия")
                self.assertEqual(items[0]["label"], "2026-11 — Гарантия")
                meta = activate_workbook_name(path.name)
                self.assertEqual(meta["name"], path.name)
                self.assertEqual(meta["year_month"], "2026-11")
                with self.assertRaises(WorkbookExistsError):
                    create_month_workbook(2026, 11, activate=False, title="Гарантия")
                other = create_month_workbook(2026, 11, activate=False, title="Луховицы")
                self.assertEqual(other.name, "Луховицы_2026-11.xlsx")
                self.assertEqual(len(list_month_workbooks()), 2)
                # Legacy SE_Lukhovitsy prefix stripped from user input
                self.assertEqual(sanitize_workbook_title("SE_Lukhovitsy_Север"), "Север")
            finally:
                config.WORKBOOKS_DIR = old_dir
                config.ACTIVE_WORKBOOK_META = old_meta
                config.SOURCE_WORKBOOK = old_src

    def test_export_workbook_period_drops_service_sheets(self):
        from core.workbook_factory import export_workbook_period

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            old_dir = config.WORKBOOKS_DIR
            old_meta = config.ACTIVE_WORKBOOK_META
            old_src = config.SOURCE_WORKBOOK
            config.WORKBOOKS_DIR = tmp_path / "workbooks"
            config.WORKBOOKS_DIR.mkdir()
            config.ACTIVE_WORKBOOK_META = tmp_path / "active_workbook.json"
            src = tmp_path / "template.xlsx"
            wb = openpyxl.Workbook()
            wb.active.title = "Данные"
            wb.create_sheet("Расчет")
            wb.create_sheet("Шаблон")
            for day in range(1, 32):
                wb.create_sheet(f"{day:02d}")
            wb.save(src)
            wb.close()
            config.SOURCE_WORKBOOK = src
            try:
                path = create_month_workbook(2026, 10, activate=True, title="Гарантия")
                out = export_workbook_period(
                    path,
                    date(2026, 10, 5),
                    date(2026, 10, 7),
                    dest=tmp_path / "out.xlsx",
                )
                exported = openpyxl.load_workbook(out)
                try:
                    names = exported.sheetnames
                    self.assertEqual(names, ["05", "06", "07"])
                    self.assertNotIn("Данные", names)
                    self.assertNotIn("Расчет", names)
                    self.assertNotIn("Шаблон", names)
                finally:
                    exported.close()

                # Inclusive range 27..29 must yield exactly three day sheets
                out29 = export_workbook_period(
                    path,
                    date(2026, 10, 27),
                    date(2026, 10, 29),
                    dest=tmp_path / "out29.xlsx",
                )
                exported29 = openpyxl.load_workbook(out29)
                try:
                    self.assertEqual(exported29.sheetnames, ["27", "28", "29"])
                finally:
                    exported29.close()

                with self.assertRaises(ValueError):
                    export_workbook_period(path, date(2026, 9, 1), date(2026, 9, 5))
                with self.assertRaises(ValueError):
                    export_workbook_period(path, date(2026, 10, 10), date(2026, 10, 5))
            finally:
                config.WORKBOOKS_DIR = old_dir
                config.ACTIVE_WORKBOOK_META = old_meta
                config.SOURCE_WORKBOOK = old_src

    def test_write_uses_titled_book_not_bare_month(self):
        """Report for YYYY-MM must go into Title_YYYY-MM.xlsx, not create bare YYYY-MM.xlsx."""
        self.assertTrue(config.SOURCE_WORKBOOK.exists())
        from core.excel_writer import append_calc_to_workbook
        from core.pricing import calculate_report
        from core.workbook_factory import ensure_month_workbook, get_active_meta, month_workbook_path

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            old_dir = config.WORKBOOKS_DIR
            old_meta = config.ACTIVE_WORKBOOK_META
            config.WORKBOOKS_DIR = tmp_path / "workbooks"
            config.WORKBOOKS_DIR.mkdir()
            config.ACTIVE_WORKBOOK_META = tmp_path / "active_workbook.json"
            try:
                titled = create_month_workbook(2026, 10, activate=True, title="Гарантия")
                self.assertEqual(titled.name, "Гарантия_2026-10.xlsx")

                path, created = ensure_month_workbook(2026, 10, activate=True)
                self.assertFalse(created)
                self.assertEqual(path.resolve(), titled.resolve())
                self.assertFalse(month_workbook_path(2026, 10).exists())

                report = {
                    "report_date": "2026-10-15",
                    "source_text_hash": "titled-book-hash",
                    "workers": [
                        {
                            "raw_worker_name": "Карандак В.Е.",
                            "members": ["Карандак В.Е."],
                            "performed_jobs": [
                                {
                                    "raw_task_name": "обжим рж45",
                                    "job_description": "обжим рж45",
                                    "service_type": "Интернет",
                                    "volume": 1,
                                    "unit": "шт",
                                }
                            ],
                        }
                    ],
                }
                calc = calculate_report(report)
                result = append_calc_to_workbook(calc, skip_duplicates=True)
                self.assertTrue(result["ok"])
                self.assertFalse(result["workbook_created"])
                self.assertEqual(Path(result["workbook"]).resolve(), titled.resolve())
                self.assertEqual(result["year_month"], "2026-10")
                self.assertFalse(month_workbook_path(2026, 10).exists())
                self.assertEqual(get_active_meta()["name"], "Гарантия_2026-10.xlsx")
            finally:
                config.WORKBOOKS_DIR = old_dir
                config.ACTIVE_WORKBOOK_META = old_meta


if __name__ == "__main__":
    unittest.main()
