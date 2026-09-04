# -*- coding: utf-8 -*-
"""Price matching and wage calculation from Данные sheet / cache."""
from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from core import config
from core.parser import resolve_members
from core.parser import strip_volume_from_task_fields


def normalize_text(value: str) -> str:
    text = unicodedata.normalize("NFKC", value or "")
    text = text.lower().replace("ё", "е")
    text = text.replace("–", "-").replace("—", "-")
    text = re.sub(r"[«»\"']", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"[.,;:]+$", "", text)
    return text

# (правка Новиков С.С. 29.07.2026)
_STOP_WORDS = {
    "монтаж", "демонтаж", "работы", "работа", "по", "для", "и", "с", "на", "до",
    "от", "до", "бр", "бр.", "заявка", "заявки"
}


def normalize_tokens(value: str) -> set[str]:
    """
    Нормализовать строку в набор значимых токенов (для сравнения по словам).

    Убираем стоп-слова, служебные слова и очень короткие токены.
    """
    base = normalize_text(value)
    tokens = re.findall(r"[a-zа-я0-9]+", base, flags=re.IGNORECASE)
    out: set[str] = set()
    for t in tokens:
        if len(t) <= 2:
            continue
        if t in _STOP_WORDS:
            continue
        out.add(t)
    return out


def load_job_aliases(path: Path | None = None) -> dict[str, str]:
    """
    Load synonym → official price name map from data/job_aliases.json.

    Supported formats:
    - { "aliases": [ { "official": "...", "synonyms": ["...", ...] }, ... ] }
    - flat { "синоним": "Официальное имя", ... } (legacy-friendly)
    """
    path = path or config.JOB_ALIASES_CACHE
    if not path.exists():
        return {}
    raw = json.loads(path.read_text(encoding="utf-8"))
    mapping: dict[str, str] = {}

    if isinstance(raw, dict) and isinstance(raw.get("aliases"), list):
        for block in raw["aliases"]:
            if not isinstance(block, dict):
                continue
            official = (block.get("official") or "").strip()
            if not official:
                continue
            synonyms = block.get("synonyms") or []
            if not isinstance(synonyms, list):
                continue
            for syn in synonyms:
                key = normalize_text(str(syn))
                if key:
                    mapping[key] = official
            # official name also matches itself
            mapping[normalize_text(official)] = official
        return mapping

    if isinstance(raw, dict):
        for syn, official in raw.items():
            if str(syn).startswith("_"):
                continue
            if not isinstance(official, str) or not official.strip():
                continue
            key = normalize_text(str(syn))
            if key:
                mapping[key] = official.strip()
    return mapping


_ALIASES_MTIME: float | None = None
MANUAL_ALIASES: dict[str, str] = {}


def get_job_aliases(path: Path | None = None) -> dict[str, str]:
    """Return alias map; reloads file when it changes on disk."""
    global MANUAL_ALIASES, _ALIASES_MTIME
    path = path or config.JOB_ALIASES_CACHE
    try:
        mtime = path.stat().st_mtime if path.exists() else None
    except OSError:
        mtime = None
    if mtime != _ALIASES_MTIME or not MANUAL_ALIASES:
        MANUAL_ALIASES = load_job_aliases(path)
        _ALIASES_MTIME = mtime
    return MANUAL_ALIASES


def reload_job_aliases(path: Path | None = None) -> dict[str, str]:
    global _ALIASES_MTIME
    _ALIASES_MTIME = None
    return get_job_aliases(path)


@dataclass
class PriceItem:
    name: str
    price: float
    unit: str = ""
    row: int | None = None
    # (правка Новиков С.С. 29.07.2026)
    service_type: str = "Общее"


@dataclass
class PricedJob:
    raw_task_name: str
    job_description: str
    matched_name: str | None
    service_type: str
    volume: float
    unit: str
    unit_price: float | None
    amount: float | None
    matched: bool
    match_score: float = 0.0


@dataclass
class WorkerCalc:
    raw_worker_name: str
    members: list[str] = field(default_factory=list)
    jobs: list[PricedJob] = field(default_factory=list)
    total: float = 0.0
    per_person: float = 0.0
    unmatched: list[PricedJob] = field(default_factory=list)

    @property
    def headcount(self) -> int:
        return max(1, len(self.members) if self.members else 1)

# исправление для подсчета Новиков С.С. (28.07)
def aggregate_by_matched_name(jobs: list[PricedJob]) -> list[PricedJob]:
    """
    Универсальная агрегация: для всех работ с одинаковым matched_name,
    unit и unit_price складывает volume и amount в одну запись.
    Остальные работы остаются как есть.
    """
    grouped: dict[tuple[str, str, float], list[PricedJob]] = {}
    order: list[tuple[str, str, float]] = []

    for job in jobs:
        if not job.matched or not job.matched_name or job.unit_price is None:
            # несопоставленные или без цены не агрегируем
            key = ("__unmatched__", job.raw_task_name, 0.0)
            if key not in grouped:
                grouped[key] = []
                order.append(key)
            grouped[key].append(job)
            continue

        key = (job.matched_name, job.unit or "", float(job.unit_price))
        if key not in grouped:
            grouped[key] = []
            order.append(key)
        grouped[key].append(job)

    aggregated: list[PricedJob] = []

    for key in order:
        jobs_in_group = grouped[key]
        # спец-ключ для unmatched — просто переносим как есть
        if key[0] == "__unmatched__":
            aggregated.extend(jobs_in_group)
            continue

        # если в группе одна работа, просто добавляем её
        if len(jobs_in_group) == 1:
            aggregated.extend(jobs_in_group)
            continue

        # несколько работ с одинаковым matched_name/unit/unit_price → объединяем
        base = jobs_in_group[0]
        total_volume = 0.0
        total_amount = 0.0
        for job in jobs_in_group:
            try:
                v = float(job.volume or 0)
            except (TypeError, ValueError):
                v = 0.0
            total_volume += v
            total_amount += float(job.amount or 0.0)

        merged = PricedJob(
            raw_task_name=base.raw_task_name,
            job_description=base.job_description,
            matched_name=base.matched_name,
            service_type=base.service_type,
            volume=total_volume,
            unit=base.unit,
            unit_price=float(base.unit_price or 0.0),
            amount=round(total_amount, 2),
            matched=True,
            match_score=base.match_score,
        )
        aggregated.append(merged)

    return aggregated

@dataclass
class CalcResult:
    report_date: str | None
    source_text_hash: str | None
    workers: list[WorkerCalc]
    grand_total: float
    unmatched_count: int

    def to_dict(self) -> dict[str, Any]:
        return {
            "report_date": self.report_date,
            "source_text_hash": self.source_text_hash,
            "grand_total": self.grand_total,
            "unmatched_count": self.unmatched_count,
            "workers": [
                {
                    "raw_worker_name": w.raw_worker_name,
                    "members": w.members,
                    "headcount": w.headcount,
                    "total": w.total,
                    "per_person": w.per_person,
                    "jobs": [asdict(j) for j in w.jobs],
                    "unmatched": [asdict(j) for j in w.unmatched],
                }
                for w in self.workers
            ],
        }


def _service_compatible(item_service: str, job_service: str) -> bool:
    """«Общее» in price or job means no hard filter."""
    item_st = (item_service or "Общее").strip() or "Общее"
    job_st = (job_service or "Общее").strip() or "Общее"
    if item_st == "Общее" or job_st == "Общее":
        return True
    return item_st == job_st


class PriceList:
    def __init__(self, items: list[PriceItem]):
        self.items = items
        self._by_norm = {normalize_text(i.name): i for i in items}

    @classmethod
    def from_cache(cls, path: Path | None = None) -> "PriceList":
        path = path or config.PRICE_CACHE
        raw = json.loads(path.read_text(encoding="utf-8"))
        items = [
            PriceItem(
                name=x["name"],
                price=float(x["price"]),
                unit=x.get("unit") or "",
                row=x.get("row"),
                # (правка Новиков С.С. 29.07.2026) тип услуги, если задан в кэше
                service_type=str(x.get("service_type") or "Общее").strip() or "Общее",
            )
            for x in raw
        ]
        return cls(items)

    def to_cache_list(self) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for i, item in enumerate(self.items):
            row: dict[str, Any] = {
                "name": item.name,
                "price": item.price,
                "unit": item.unit,
                "row": item.row if item.row is not None else i + config.PRICE_START_ROW,
            }
            if item.service_type and item.service_type != "Общее":
                row["service_type"] = item.service_type
            out.append(row)
        return out

    def save_cache(self, path: Path | None = None) -> Path:
        path = path or config.PRICE_CACHE
        path.write_text(
            json.dumps(self.to_cache_list(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return path

    @classmethod
    def from_workbook(cls, workbook_path: Path | None = None) -> "PriceList":
        import openpyxl

        path = workbook_path or config.WORKING_WORKBOOK
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[config.SHEET_DATA]
        items: list[PriceItem] = []
        for r in range(config.PRICE_START_ROW, 400):
            name = ws.cell(r, config.PRICE_NAME_COL).value
            price = ws.cell(r, config.PRICE_VALUE_COL).value
            unit = ws.cell(r, config.PRICE_UNIT_COL).value
            if name and isinstance(price, (int, float)):
                items.append(
                    PriceItem(
                        name=str(name).strip(),
                        price=float(price),
                        unit=str(unit or "").strip(),
                        row=r,
                    )
                )
        wb.close()
        return cls(items)

    def names(self) -> list[str]:
        return [i.name for i in self.items]

    def match(
            self,
            description: str,
            raw_task: str = "",
            service_type: str = "Общее",
    ) -> tuple[PriceItem | None, float]:
        service_type = (service_type or "Общее").strip()
        candidates = [description or "", raw_task or ""]
        aliases = get_job_aliases()

        # 1) Пробуем алиасы и точное совпадение в рамках service_type
        for cand in candidates:
            norm = normalize_text(cand)
            if not norm:
                continue

            # алиас → official
            alias = aliases.get(norm)
            if alias:
                item = self._by_norm.get(normalize_text(alias))
                if item and _service_compatible(item.service_type, service_type):
                    return item, 1.0

            # прямое совпадение с позицией прайса
            item = self._by_norm.get(norm)
            if item and _service_compatible(item.service_type, service_type):
                return item, 1.0

        # 2) Нечёткий поиск по прайсу (сначала в своём service_type)
        best: PriceItem | None = None
        best_score = 0.0

        def consider_item(it: PriceItem, cand_norm: str) -> None:
            nonlocal best, best_score
            inorm = normalize_text(it.name)
            score = _similarity(cand_norm, inorm)
            if score > best_score:
                best_score = score
                best = it

        # 2а) сначала позиции с совместимым service_type
        for cand in candidates:
            norm = normalize_text(cand)
            if len(norm) < 4:
                continue
            for item in self.items:
                if not _service_compatible(item.service_type, service_type):
                    continue
                consider_item(item, norm)

        # 2б) если ничего хорошего не нашли — смотрим на весь прайс
        if best is None or best_score < 0.72:
            for cand in candidates:
                norm = normalize_text(cand)
                if len(norm) < 4:
                    continue
                for item in self.items:
                    consider_item(item, norm)

        if best and best_score >= 0.72:
            return best, best_score
        return None, best_score


def save_price_items(items: list[dict[str, Any]], path: Path | None = None) -> list[dict[str, Any]]:
    """Validate and persist price list JSON. Returns normalized items."""
    path = path or config.PRICE_CACHE
    normalized: list[PriceItem] = []
    for idx, raw in enumerate(items):
        if not isinstance(raw, dict):
            raise ValueError(f"Строка {idx + 1}: ожидается объект")
        name = str(raw.get("name") or "").strip()
        if not name:
            raise ValueError(f"Строка {idx + 1}: пустое наименование")
        try:
            price = float(raw.get("price"))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Строка {idx + 1}: некорректная цена") from exc
        unit = str(raw.get("unit") or "").strip()
        row = raw.get("row")
        if row is not None:
            try:
                row = int(row)
            except (TypeError, ValueError):
                row = None
        service_type = str(raw.get("service_type") or "Общее").strip() or "Общее"
        normalized.append(
            PriceItem(name=name, price=price, unit=unit, row=row, service_type=service_type)
        )
    pl = PriceList(normalized)
    pl.save_cache(path)
    return pl.to_cache_list()


def load_aliases_document(path: Path | None = None) -> dict[str, Any]:
    path = path or config.JOB_ALIASES_CACHE
    if not path.exists():
        return {"aliases": []}
    raw = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(raw, dict) and isinstance(raw.get("aliases"), list):
        return raw
    # legacy flat map → structured
    aliases = []
    if isinstance(raw, dict):
        by_official: dict[str, list[str]] = {}
        for syn, official in raw.items():
            if str(syn).startswith("_") or not isinstance(official, str):
                continue
            by_official.setdefault(official.strip(), []).append(str(syn))
        aliases = [{"official": k, "synonyms": v} for k, v in by_official.items()]
    return {"aliases": aliases}


def save_aliases_document(doc: dict[str, Any], path: Path | None = None) -> dict[str, Any]:
    path = path or config.JOB_ALIASES_CACHE
    if not isinstance(doc, dict):
        raise ValueError("Ожидается объект с полем aliases")
    aliases = doc.get("aliases")
    if not isinstance(aliases, list):
        raise ValueError("Поле aliases должно быть массивом")
    cleaned = []
    for idx, block in enumerate(aliases):
        if not isinstance(block, dict):
            raise ValueError(f"Алиас {idx + 1}: ожидается объект")
        official = str(block.get("official") or "").strip()
        if not official:
            raise ValueError(f"Алиас {idx + 1}: пустое official")
        synonyms = block.get("synonyms") or []
        if not isinstance(synonyms, list):
            raise ValueError(f"Алиас {idx + 1}: synonyms должен быть массивом")
        cleaned.append(
            {
                "official": official,
                "synonyms": [str(s).strip() for s in synonyms if str(s).strip()],
            }
        )
    out = {k: v for k, v in doc.items() if k != "aliases"}
    out["aliases"] = cleaned
    path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    reload_job_aliases(path)
    return out


def _similarity(a: str, b: str) -> float:
    # (правка Новиков С.С. 29.07.2026)
    # Используем два уровня:
    # 1) быстрое сравнение нормализованных строк;
    # 2) сравнение наборов токенов с Jaccard.
    if not a or not b:
        return 0.0

    na = normalize_text(a)
    nb = normalize_text(b)
    if not na or not nb:
        return 0.0

    if na == nb:
        return 1.0
    if na in nb or nb in na:
        shorter, longer = (na, nb) if len(na) <= len(nb) else (nb, na)
        return 0.75 + 0.25 * (len(shorter) / len(longer))

    ta = normalize_tokens(na)
    tb = normalize_tokens(nb)
    if not ta or not tb:
        return 0.0
    inter = len(ta & tb)
    union = len(ta | tb)
    jacc = inter / union

    # чуть усиливаем, если совпали ключевые слова
    if inter >= 2:
        jacc += 0.1
    if jacc > 1.0:
        jacc = 1.0
    return jacc

# правки в коде Новиков С.С. (28.07.2026)
def calculate_report(report: dict[str, Any], price_list: PriceList | None = None) -> CalcResult:
    pl = price_list or PriceList.from_cache()
    workers_out: list[WorkerCalc] = []
    unmatched_count = 0
    grand = 0.0

    for w in report.get("workers") or []:
        members = w.get("members") or []
        if not isinstance(members, list) or not members:
            name = (w.get("raw_worker_name") or "").strip()
            members = [name] if name else []
        members = resolve_members([str(m).strip() for m in members if str(m).strip()])
        display = ", ".join(members) if members else (w.get("raw_worker_name") or "").strip()
        wc = WorkerCalc(raw_worker_name=display, members=members)

        for j in w.get("performed_jobs") or []:
            raw_name = str(j.get("raw_task_name") or "")
            desc = str(j.get("job_description") or "")
            volume = float(j.get("volume") or 0)
            unit = str(j.get("unit") or "")

            # убираем хвосты с количеством и «шт/м», чтобы алиасы работали по чистому тексту
            clean_raw, clean_desc, _, _ = strip_volume_from_task_fields(
                raw_name,
                desc,
                volume,
                unit,
            )

            service_type = j.get("service_type") or "Общее"

            # в match() передаём уже очищенный текст и тип услуги
            item, score = pl.match(
                clean_desc or desc,
                clean_raw or raw_name,
                service_type=service_type,
            )

            if item:
                amount = round(volume * item.price, 2)
                pj = PricedJob(
                    raw_task_name=raw_name,
                    job_description=desc,
                    matched_name=item.name,
                    service_type=service_type,
                    volume=volume,
                    unit=unit or item.unit,
                    unit_price=item.price,
                    amount=amount,
                    matched=True,
                    match_score=score,
                )
                wc.jobs.append(pj)
                wc.total += amount
            else:
                pj = PricedJob(
                    raw_task_name=raw_name,
                    job_description=desc,
                    matched_name=None,
                    service_type=service_type,
                    volume=volume,
                    unit=unit,
                    unit_price=None,
                    amount=None,
                    matched=False,
                    match_score=score,
                )
                wc.jobs.append(pj)
                wc.unmatched.append(pj)
                unmatched_count += 1

        # АГРЕГАЦИЯ: слить работы с одинаковым matched_name/unit/unit_price
        wc.jobs = aggregate_by_matched_name(wc.jobs)

        # Пересчитать total и per_person уже по агрегированным jobs
        wc.total = round(sum(job.amount or 0.0 for job in wc.jobs), 2)
        wc.per_person = round(wc.total / wc.headcount, 2)
        grand += wc.total
        workers_out.append(wc)

    return CalcResult(
        report_date=report.get("report_date"),
        source_text_hash=report.get("source_text_hash"),
        workers=workers_out,
        grand_total=round(grand, 2),
        unmatched_count=unmatched_count,
    )
