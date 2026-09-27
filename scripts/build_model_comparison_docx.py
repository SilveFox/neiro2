# -*- coding: utf-8 -*-
"""Generate comparative AI model report (Word) for MontazhPro."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "docs" / "Sravnenie_modelej_II_MontazhPro.docx"


def set_run_font(run, name: str = "Calibri", size: int = 11, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_ru(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "Calibri", 16 if level == 1 else 13, bold=True)
    return p


def add_para(doc: Document, text: str, *, bold: bool = False, size: int = 11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def shade_header_row(row):
    from docx.oxml import OxmlElement

    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.first_child_found_in("w:shd")
        if shd is None:
            shd = OxmlElement("w:shd")
            tcPr.append(shd)
        shd.set(qn("w:fill"), "1F54C4")
        shd.set(qn("w:val"), "clear")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True


def fill_table(table, headers: list[str], rows: list[list[str]]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True)
    shade_header_row(table.rows[0])
    for r_i, row_data in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row_data):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=9)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Сравнительный отчёт")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Выбор локальной ИИ-модели (Ollama) для приложения «МонтажПро»\n"
        "разбор сменных отчётов монтажников → JSON → расчёт ЗП → Excel"
    )
    set_run_font(r, size=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(f"Дата: {date.today().strftime('%d.%m.%Y')}  ·  Целевая среда: сервер без GPU (CPU-only)")
    set_run_font(r, size=10)
    for run in meta.runs:
        run.font.color.rgb = RGBColor(80, 90, 110)

    doc.add_paragraph()

    # 1. Purpose
    add_heading_ru(doc, "1. Цель отчёта", 1)
    add_para(
        doc,
        "Подобрать модель для локального разбора текстовых отчётов монтажников в приложении "
        "МонтажПро (Ollama + FastAPI). Модель должна стабильно выдавать JSON по схеме workers[], "
        "корректно обрабатывать русский текст (бригады, объёмы, названия работ) и работать "
        "приемлемо на CPU без видеокарты.",
    )

    # 2. Requirements
    add_heading_ru(doc, "2. Требования приложения к модели", 1)
    add_para(doc, "Обязательные:")
    for item in (
        "Русский язык: фамилии, адреса, сленг монтажных работ («сварки 1-7», «обжим RJ-45»).",
        "Структурированный JSON: report_date, workers[], members[], performed_jobs[] с volume/unit/service_type.",
        "Следование системному промпту (бригада = один worker, объём не в названии работы).",
        "Локальный запуск через Ollama без облачных API.",
        "Работа на CPU (сервер без GPU); разумное время ответа (ориентир до 1–3 минут на отчёт).",
    ):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)

    add_para(doc, "Желательные:")
    for item in (
        "Размер модели на диске ≤ 3 GB (удобно для серверов с 8–16 GB RAM).",
        "Стабильный JSON mode (format=json в Ollama).",
        "Минимум «галлюцинаций» объёмов и лишних работ из адресов/заголовков заявок.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)

    # 3. Criteria
    add_heading_ru(doc, "3. Критерии сравнения", 1)
    add_para(
        doc,
        "Оценки по 5-балльной шкале (5 — отлично подходит под задачу МонтажПро, 1 — плохо подходит). "
        "Оценки экспертные, с опорой на публичные бенчмарки малых моделей на CPU, документацию Ollama/Qwen "
        "и опыт использования qwen2.5:7b / компактного промпта в проекте.",
    )

    criteria = [
        ["Критерий", "Что измеряем"],
        ["Скорость на CPU", "Токены/с, время разбора типичного отчёта"],
        ["Русский язык", "Понимание отчётов, ФИО, единиц измерения"],
        ["JSON / схема", "Валидный JSON, полнота полей, меньше retry"],
        ["Качество разбора", "Бригады, объёмы (в т.ч. «19 сварок 1-7», «30 м»)"],
        ["Ресурсы", "Размер на диске, RAM, пригодность без GPU"],
        ["Совместимость", "Семья модели vs текущий промпт МонтажПро"],
    ]
    t = doc.add_table(rows=len(criteria), cols=2)
    t.style = "Table Grid"
    fill_table(t, criteria[0], criteria[1:])
    doc.add_paragraph()

    # 4. Models compared
    add_heading_ru(doc, "4. Сравниваемые модели", 1)
    add_para(
        doc,
        "Отобраны модели, доступные в Ollama и реалистичные для CPU-only сервера "
        "(примерно до 8B параметров). Более крупные (14B+) для этой среды не рассматриваются.",
    )

    models_overview = [
        ["Модель (Ollama)", "Параметры", "~Размер", "Роль в сравнении"],
        ["qwen2.5:7b", "7B", "~4,7 GB", "Исходная модель проекта (эталон качества)"],
        ["qwen2.5:3b", "3B", "~1,9 GB", "Основной кандидат для CPU"],
        ["qwen2.5:1.5b", "1.5B", "~1,0 GB", "Запасной вариант на слабом CPU"],
        ["qwen3:1.7b", "1.7B", "~1,4 GB", "Новое поколение Qwen; сильнее в агентных задачах"],
        ["llama3.2:3b", "3B", "~2,0 GB", "Популярная малая модель Meta"],
        ["phi3.5:3.8b", "~3.8B", "~2,2 GB", "Сильна в коде/инструкциях, слабее в русском"],
        ["gemma3:4b", "4B", "~3,1 GB", "Альтернатива Google; тяжелее на CPU"],
    ]
    t2 = doc.add_table(rows=len(models_overview), cols=4)
    t2.style = "Table Grid"
    fill_table(t2, models_overview[0], models_overview[1:])
    doc.add_paragraph()

    # 5. Score matrix
    add_heading_ru(doc, "5. Сводная матрица оценок", 1)
    scores = [
        ["Модель", "CPU", "RU", "JSON", "Разбор", "Ресурсы", "Совмест.", "Итого"],
        ["qwen2.5:7b", "2", "5", "5", "5", "2", "5", "24"],
        ["qwen2.5:3b", "4", "5", "5", "4", "5", "5", "28"],
        ["qwen2.5:1.5b", "5", "4", "4", "3", "5", "5", "26"],
        ["qwen3:1.7b", "3", "4", "4", "3", "5", "3", "22"],
        ["llama3.2:3b", "4", "3", "3", "3", "5", "2", "20"],
        ["phi3.5:3.8b", "3", "2", "4", "2", "4", "2", "17"],
        ["gemma3:4b", "3", "3", "4", "3", "3", "2", "18"],
    ]
    t3 = doc.add_table(rows=len(scores), cols=8)
    t3.style = "Table Grid"
    fill_table(t3, scores[0], scores[1:])
    add_para(
        doc,
        "Итого = сумма баллов по шести критериям (макс. 30). Лидер для CPU-only: qwen2.5:3b (28).",
        bold=False,
    )

    # 6. Detailed
    add_heading_ru(doc, "6. Разбор по моделям", 1)

    add_heading_ru(doc, "6.1. qwen2.5:7b — эталон качества", 2)
    add_para(
        doc,
        "Использовалась в проекте изначально. Хорошо держит схему JSON, русские отчёты и правила "
        "промпта (бригада, «19 сварок 1-7», длины в метрах). На CPU без GPU заметно медленнее 3B "
        "(холодный старт и генерация), требует больше RAM (~5–8+ GB под модель). "
        "Рекомендуется только при появлении GPU или очень мощном многоядерном CPU.",
    )

    add_heading_ru(doc, "6.2. qwen2.5:3b — рекомендуемая для продакшена на CPU", 2)
    add_para(
        doc,
        "Лучший баланс для МонтажПро на сервере без GPU. Та же семья, что и 7B: текущий компактный "
        "промпт почти не нужно переписывать. Официально поддерживает русский и structured JSON. "
        "В публичных CPU-бенчмарках малых моделей (~3B) часто проходит тесты tool/JSON быстрее 7B "
        "при приемлемом качестве. Размер ~2 GB. В конфиге проекта уже указана как MODEL_NAME по умолчанию.",
    )

    add_heading_ru(doc, "6.3. qwen2.5:1.5b — запасной вариант", 2)
    add_para(
        doc,
        "Самая быстрая из линейки Qwen2.5 среди разумных кандидатов. На очень слабом CPU даёт "
        "минимальную задержку, но чаще ошибается в краевых случаях: лишние работы из «1 заявка…», "
        "пропуск полей JSON, путаница объёмов. Имеет смысл как fallback (OLLAMA fallback) или "
        "для отладки пайплайна, не как основной выбор при достаточной RAM.",
    )

    add_heading_ru(doc, "6.4. qwen3:1.7b — осторожно", 2)
    add_para(
        doc,
        "Сильна в tool-calling бенчмарках, но режим «размышления» (thinking) увеличивает latency "
        "и может засорять ответ лишним текстом до JSON. Для МонтажПро критичен чистый JSON без "
        "Markdown. Потребует отдельной настройки (отключение thinking / жёсткий format=json) "
        "и повторного тестирования промпта. Пока не рекомендуется как drop-in замена.",
    )

    add_heading_ru(doc, "6.5. llama3.2:3b", 2)
    add_para(
        doc,
        "Быстрая на CPU, но слабее Qwen в русском и в дисциплине JSON-схемы для многополевого "
        "извлечения. Часто избыточно «болтлива». Для отчётов монтажников на русском — хуже посадка, "
        "чем у qwen2.5:3b при сопоставимом размере.",
    )

    add_heading_ru(doc, "6.6. phi3.5:3.8b и gemma3:4b", 2)
    add_para(
        doc,
        "Phi ориентирована на код и английские инструкции; русский и доменный сленг — слабое место. "
        "Gemma 3 4B тяжелее на CPU и хуже стыкуется с уже отлаженным Qwen-промптом. "
        "Обе модели требуют заметной переработки промпта и регрессионных тестов — выгода сомнительна.",
    )

    # 7. Fit to pipeline
    add_heading_ru(doc, "7. Соответствие пайплайну МонтажПро", 1)
    add_para(
        doc,
        "Пайплайн приложения: текст → предобработка объёмов → Ollama (system prompt) → "
        "валидация JSON → reconcile объёмов → прайс/алиасы → Excel. "
        "Модель влияет только на этап Ollama; постобработка частично компенсирует ошибки объёмов, "
        "но не спасает от битого JSON или неверной бригады. Поэтому приоритет: JSON + русский + "
        "совместимость с текущим промптом выше «сырой» скорости.",
    )

    pipeline = [
        ["Этап", "Чувствительность к модели", "Комментарий"],
        ["Предобработка текста", "Низкая", "Код Python, без LLM"],
        ["Вызов Ollama", "Высокая", "Качество JSON и семантика отчёта"],
        ["validate_report", "Средняя", "Отсекает битый JSON; retry при ошибке"],
        ["reconcile_volumes", "Средняя", "Чинит часть ошибок объёмов по исходнику"],
        ["Прайс / алиасы", "Низкая", "Не зависит от модели"],
        ["Запись Excel", "Низкая", "Не зависит от модели"],
    ]
    t4 = doc.add_table(rows=len(pipeline), cols=3)
    t4.style = "Table Grid"
    fill_table(t4, pipeline[0], pipeline[1:])
    doc.add_paragraph()

    # 8. Recommendations
    add_heading_ru(doc, "8. Рекомендации", 1)

    add_heading_ru(doc, "8.1. Итоговый выбор", 2)
    recs = [
        ["Сценарий", "Модель", "Действие"],
        ["Продакшен CPU (текущий сервер)", "qwen2.5:3b", "Оставить / установить: ollama pull qwen2.5:3b"],
        ["Очень слабый CPU / мало RAM", "qwen2.5:1.5b", "Сменить MODEL_NAME, усилить автотесты JSON"],
        ["Появился GPU (8+ GB VRAM)", "qwen2.5:7b", "Вернуть 7b для максимального качества"],
        ["Эксперименты", "qwen3:1.7b", "Только после A/B с tests/test_prompt_compact.py"],
    ]
    t5 = doc.add_table(rows=len(recs), cols=3)
    t5.style = "Table Grid"
    fill_table(t5, recs[0], recs[1:])
    doc.add_paragraph()

    add_heading_ru(doc, "8.2. Порядок внедрения на сервере", 2)
    for i, step in enumerate(
        (
            "Установить Ollama (если ещё нет).",
            "Выполнить: ollama pull qwen2.5:3b",
            "В core/config.py убедиться: MODEL_NAME = \"qwen2.5:3b\".",
            "Перезапустить веб-приложение (run_web.py / служба).",
            "Прогнать 3–5 реальных отчётов: бригада, объёмы в начале/конце строки, «сварки 1-7».",
            "При частых ошибках JSON — временно USE_FULL_PROMPT = True или вернуться к 7b на более мощной машине.",
        ),
        start=1,
    ):
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(step)
        set_run_font(run, size=11)

    # 9. Risks
    add_heading_ru(doc, "9. Риски и ограничения", 1)
    for item in (
        "Оценки в матрице — экспертные; на конкретном CPU (число ядер, AVX, RAM) скорость будет отличаться.",
        "Смена семейства модели (Llama/Gemma/Phi) потребует повторной настройки промпта и регрессии.",
        "Компактный промпт (~4× короче полного) ускоряет все модели, но мелкие модели чаще игнорируют редкие правила.",
        "Тарифы и алиасы не зависят от LLM: «красные» строки в UI чаще решаются справочниками, а не сменой модели.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)

    # 10. Conclusion
    add_heading_ru(doc, "10. Заключение", 1)
    add_para(
        doc,
        "Для сервера без GPU оптимальная модель для МонтажПро — qwen2.5:3b: сохраняет сильные стороны "
        "линейки Qwen2.5 (русский язык, JSON), укладывается в типичный объём RAM и даёт лучший "
        "суммарный балл среди рассмотренных вариантов. qwen2.5:7b остаётся эталоном качества при GPU; "
        "qwen2.5:1.5b — резерв для слабого железа. Смена на Llama/Gemma/Phi для данной задачи "
        "нецелесообразна без отдельного проекта адаптации промпта.",
    )

    add_para(
        doc,
        "Документ подготовлен для внутреннего использования команды МонтажПро. "
        "Файл: docs/Sravnenie_modelej_II_MontazhPro.docx",
        size=10,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
