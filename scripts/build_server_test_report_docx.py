# -*- coding: utf-8 -*-
"""Generate server testing report (Word) for MontazhPro."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "docs" / "Otchet_testirovanie_server_MontazhPro.docx"


def set_run_font(run, name: str = "Calibri", size: int = 11, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_ru(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "Calibri", 16 if level == 1 else 13, bold=True)
    return p


def add_para(doc: Document, text: str, *, bold: bool = False, size: int = 11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def shade_header_row(row):
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = tcPr.first_child_found_in("w:shd")
        if shd is None:
            shd = OxmlElement("w:shd")
            tcPr.append(shd)
        shd.set(qn("w:fill"), "1F54C4")
        shd.set(qn("w:val"), "clear")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True


def fill_table(table, headers: list[str], rows: list[list[str]]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True)
    shade_header_row(table.rows[0])
    for r_i, row_data in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row_data):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, size=9)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Отчёт по тестированию")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(
        "Приложение «МонтажПро» на рабочем сервере\n"
        "разбор отчётов → прайс / алиасы → расчёт ЗП → Excel"
    )
    set_run_font(r, size=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Дата: {date.today().strftime('%d.%m.%Y')}  ·  "
        "Среда: Linux-сервер /opt/neiro, без GPU (CPU-only)"
    )
    set_run_font(r, size=10)
    for run in meta.runs:
        run.font.color.rgb = RGBColor(80, 90, 110)

    doc.add_paragraph()

    # 1
    add_heading_ru(doc, "1. Цель тестирования", 1)
    add_para(
        doc,
        "Проверить работоспособность МонтажПро после развёртывания на сервере: "
        "разбор текстовых отчётов через Ollama, штатную работу справочников "
        "«Прайс» и «Алиасы работ», расчёт заработной платы и запись/скачивание Excel. "
        "Зафиксировать ограничения по скорости на CPU без видеокарты.",
    )

    # 2
    add_heading_ru(doc, "2. Тестовая среда", 1)
    env = [
        ["Параметр", "Значение"],
        ["Каталог на сервере", "/opt/neiro"],
        ["ОС", "Linux (python3)"],
        ["Доступ", "SSH"],
        ["Репозиторий", "https://github.com/SilveFox/neiro2.git"],
        ["Запуск", "python3 run_web.py"],
        ["Интерфейс", "http://<IP>:8000"],
        ["ИИ-модель (Ollama)", "qwen2.5:3b (ориентир для CPU)"],
        ["GPU", "Отсутствует — вывод только на CPU"],
        ["Таймаут Ollama", "OLLAMA_TIMEOUT_SEC в core/config.py (настраивается)"],
    ]
    t = doc.add_table(rows=len(env), cols=2)
    t.style = "Table Grid"
    fill_table(t, env[0], env[1:])
    doc.add_paragraph()

    # 3
    add_heading_ru(doc, "3. Объект тестирования", 1)
    add_bullet(doc, "Вкладка «Отчёт»: загрузка текста/файла, разбор через ИИ, предпросмотр, запись в Excel.")
    add_bullet(doc, "Вкладка «Прайс»: просмотр, поиск, добавление/правка, сохранение в data/price_list.json.")
    add_bullet(doc, "Вкладка «Алиасы работ»: синонимы → официальное имя прайса, поиск, сохранение.")
    add_bullet(doc, "Вкладки «Настройки», «Лог Excel», «Инструкция».")
    add_bullet(doc, "Скачивание месячной книги SE_Lukhovitsy_ГГГГ-ММ.xlsx.")

    # 4 Price
    add_heading_ru(doc, "4. Штатная работа прайса", 1)
    add_para(
        doc,
        "Прайс — основной справочник тарифов. Расчёт в интерфейсе идёт по "
        "data/price_list.json. При записи в Excel тариф и сумма сохраняются "
        "числами из того же расчёта (чтобы в скачанном файле сразу были суммы).",
    )
    add_heading_ru(doc, "4.1. Проверенный сценарий", 2)
    for s in (
        "Открыть вкладку «Прайс» — список наименований, цен, единиц и типов загружается.",
        "Поиск по наименованию отфильтровывает строки.",
        "Добавление новой строки → «Сохранить» — данные пишутся в price_list.json.",
        "Правка цены существующей работы → сохранение → повторный «Пересчитать по прайсу» "
        "в отчёте отражает новую цену в предпросмотре.",
        "Тип «Общее» подходит к работам Интернет/ТВ без жёсткого фильтра по типу услуги.",
        "Работа, отсутствующая в прайсе, в предпросмотре подсвечивается (красная строка) "
        "и не входит в денежный итог — штатное поведение.",
    ):
        add_bullet(doc, s)

    add_heading_ru(doc, "4.2. Результат", 2)
    add_para(
        doc,
        "Штатная работа прайса подтверждена: CRUD, поиск, влияние на расчёт ЗП "
        "и отображение ненайденных работ работают ожидаемо. "
        "Рекомендация: держать price_list.json согласованным с листом «Данные» в шаблоне Excel.",
        bold=False,
    )

    # 5 Aliases
    add_heading_ru(doc, "5. Штатная работа алиасов", 1)
    add_para(
        doc,
        "Алиасы связывают формулировки из живых отчётов монтажников с официальными "
        "именами прайса. Файл: data/job_aliases.json. Порядок сопоставления: "
        "синоним → точное имя → нечёткий поиск.",
    )
    add_heading_ru(doc, "5.1. Проверенный сценарий", 2)
    for s in (
        "Вкладка «Алиасы работ» открывается, список официальных имён и синонимов загружается.",
        "Поиск работает и по официальному названию, и по синонимам.",
        "Добавление синонима (через «;» несколько вариантов) → «Сохранить».",
        "Повторный разбор/пересчёт отчёта: формулировка из текста матчится на имя прайса, "
        "строка перестаёт быть «проблемной», тариф подставляется.",
        "Удаление или правка алиаса влияет на последующие расчёты после сохранения.",
    ):
        add_bullet(doc, s)

    add_heading_ru(doc, "5.2. Результат", 2)
    add_para(
        doc,
        "Штатная работа алиасов подтверждена. При появлении новых формулировок в отчётах "
        "достаточно добавить синоним — без изменения кода и без правки каждой строки вручную в JSON отчёта.",
    )

    # 6 Performance - KEY USER REQUEST
    add_heading_ru(doc, "6. Производительность разбора на CPU (без GPU)", 1)
    add_para(
        doc,
        "Сервер не оснащён GPU. Вывод модели Ollama выполняется только на процессоре. "
        "Это критично влияет на время разбора отчёта.",
    )

    add_heading_ru(doc, "6.1. Зафиксированное ограничение", 2)
    add_para(
        doc,
        "Отчёт, содержащий более 4 строк работ, на данном сервере "
        "не обрабатывается быстрее чем за 5 минут из‑за отсутствия GPU.",
        bold=True,
    )
    add_para(
        doc,
        "При увеличении числа строк (работ) растут объём промпта, длина JSON-ответа модели "
        "и время генерации на CPU. Таймаут приложения (OLLAMA_TIMEOUT_SEC) при необходимости "
        "увеличивают, но ускорить сам вывод модели без GPU нельзя.",
    )

    perf = [
        ["Условие", "Ожидаемое поведение на CPU"],
        ["До ~4 строк работ", "Разбор возможен за приемлемое время (зависит от загрузки CPU)"],
        ["Более 4 строк работ", "Время разбора ≥ 5 минут; ускорение без GPU недоступно"],
        ["Длинный отчёт / бригада + много позиций", "Риск таймаута; нужна настройка OLLAMA_TIMEOUT_SEC"],
        ["Модель qwen2.5:3b", "Компромисс скорость/качество для CPU; 7B ещё медленнее без GPU"],
    ]
    t2 = doc.add_table(rows=len(perf), cols=2)
    t2.style = "Table Grid"
    fill_table(t2, perf[0], perf[1:])
    doc.add_paragraph()

    add_heading_ru(doc, "6.2. Рекомендации по эксплуатации", 2)
    for s in (
        "Дробить крупные сменные отчёты на части (до 4 работ за один разбор) либо принимать ожидание ≥ 5 минут.",
        "Не возвращать qwen2.5:7b на этот сервер без GPU — будет ещё медленнее.",
        "При появлении GPU рассмотреть возврат к более крупной модели для качества и скорости.",
        "Таймаут в config.py выставлять с запасом относительно реального времени на CPU.",
    ):
        add_bullet(doc, s)

    # 7 Other functional
    add_heading_ru(doc, "7. Прочие проверки на сервере", 1)
    other = [
        ["Проверка", "Результат"],
        ["Обновление кода git pull origin main", "Успешно; подтянуты excel_writer, тесты, документация"],
        ["Запуск python3 run_web.py", "Рабочий способ запуска (команда python на сервере отсутствует)"],
        ["Разбор короткого отчёта (≤4 работ)", "JSON, предпросмотр, расчёт ЗП"],
        ["Запись в Excel / скачивание", "После исправления: тариф и ЗП сохраняются числами из UI"],
        ["Прайс (штатный режим)", "Работает — см. раздел 4"],
        ["Алиасы (штатный режим)", "Работают — см. раздел 5"],
        ["Производительность >4 работ", "≥ 5 мин на CPU без GPU — см. раздел 6"],
    ]
    t3 = doc.add_table(rows=len(other), cols=2)
    t3.style = "Table Grid"
    fill_table(t3, other[0], other[1:])
    doc.add_paragraph()

    # 8 Conclusions
    add_heading_ru(doc, "8. Выводы", 1)
    add_bullet(
        doc,
        "Приложение на сервере /opt/neiro работоспособно при запуске через python3.",
    )
    add_bullet(
        doc,
        "Прайс и алиасы функционируют в штатном режиме: редактирование, поиск, влияние на расчёт подтверждены.",
    )
    add_bullet(
        doc,
        "Главное эксплуатационное ограничение: без GPU отчёты с более чем 4 строками работ "
        "не разбираются быстрее 5 минут.",
    )
    add_bullet(
        doc,
        "Для ускорения разбора крупных отчётов необходимы GPU либо дробление входных данных.",
    )

    add_heading_ru(doc, "9. Статус", 1)
    add_para(
        doc,
        "Тестирование на сервере: функциональные модули прайса и алиасов — пригодны к штатной эксплуатации. "
        "Производительность ИИ-разбора на CPU — с указанным ограничением по времени.",
        bold=True,
    )

    add_para(
        doc,
        "Документ: docs/Otchet_testirovanie_server_MontazhPro.docx",
        size=10,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
