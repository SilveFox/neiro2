# -*- coding: utf-8 -*-
"""E2E test: parse sample report → calc → save to SQLite → Word test report."""
from __future__ import annotations

import json
import sys
import time
import traceback
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core import config
from core.parser import parse_report_text
from core.pricing import calculate_report
from core.report_db import get_report_details, save_report

SAMPLE_TEXT = """Карандак, Лошкомойников
Заявки по неисправности интернета 7шт
Обмер 2шт
Подключение мкд 1шт

Монтаж кабеля по стояку 15м
Монтаж кабеля по бетонной стене с креплением 25м
Сверление отверстия в бетоне до 24×500 1шт
Обжим рж45 2шт
Настройка соединения на пк 1шт
"""

EXPECT_MEMBERS_HINTS = ("карандак", "лошкомойников")
EXPECT_VOLUMES = {7, 2, 1, 15, 25, 1, 2, 1}  # soft: key volumes present
EXPECT_MIN_JOBS = 7

OUT_DOCX = ROOT / "docs" / "Otchet_test_sohranenie_BD.docx"
OUT_JSON = ROOT / "docs" / "_test_save_db_result.json"
TEST_DB = ROOT / "data" / "reports_test_save.db"


def set_run_font(run, name: str = "Calibri", size: int = 11, bold: bool = False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, "Calibri", 16 if level == 1 else 13, bold=True)


def add_para(doc, text, *, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(6)


def shade_header(row):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.first_child_found_in("w:shd")
        if shd is None:
            shd = OxmlElement("w:shd")
            tcPr.append(shd)
        shd.set(qn("w:fill"), "1F54C4")
        shd.set(qn("w:val"), "clear")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True


def fill_table(table, headers, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = ""
        run = table.rows[0].cells[i].paragraphs[0].add_run(h)
        set_run_font(run, size=9, bold=True)
    shade_header(table.rows[0])
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            table.rows[r_i + 1].cells[c_i].text = ""
            run = table.rows[r_i + 1].cells[c_i].paragraphs[0].add_run(str(val))
            set_run_font(run, size=9)


def run_pipeline() -> dict:
    results = {
        "sample": SAMPLE_TEXT,
        "steps": [],
        "checks": [],
        "ok": True,
    }

    # 1) Parse
    t0 = time.perf_counter()
    try:
        report = parse_report_text(
            SAMPLE_TEXT,
            report_date="2026-10-04",
            timeout=config.OLLAMA_TIMEOUT_SEC,
        )
        elapsed = round(time.perf_counter() - t0, 2)
        results["steps"].append(
            {
                "step": "parse",
                "ok": True,
                "elapsed_sec": elapsed,
                "workers": len(report.get("workers") or []),
            }
        )
        results["report"] = report
    except Exception as exc:
        results["steps"].append(
            {
                "step": "parse",
                "ok": False,
                "elapsed_sec": round(time.perf_counter() - t0, 2),
                "error": str(exc),
                "trace": traceback.format_exc()[-800:],
            }
        )
        results["ok"] = False
        return results

    # 2) Calc
    t1 = time.perf_counter()
    try:
        calc = calculate_report(report)
        calc_dict = calc.to_dict()
        results["steps"].append(
            {
                "step": "calculate",
                "ok": True,
                "elapsed_sec": round(time.perf_counter() - t1, 2),
                "grand_total": calc.grand_total,
                "unmatched_count": calc.unmatched_count,
            }
        )
        results["calculation"] = calc_dict
    except Exception as exc:
        results["steps"].append(
            {
                "step": "calculate",
                "ok": False,
                "error": str(exc),
            }
        )
        results["ok"] = False
        return results

    # 3) Save DB (isolated test file)
    config.REPORTS_DB = TEST_DB
    if TEST_DB.exists():
        TEST_DB.unlink()
    t2 = time.perf_counter()
    try:
        saved = save_report(report, calculation=calc_dict, source_text=SAMPLE_TEXT)
        details = get_report_details(saved["id"])
        results["steps"].append(
            {
                "step": "save_db",
                "ok": True,
                "elapsed_sec": round(time.perf_counter() - t2, 2),
                "id": saved["id"],
                "db_path": saved["db_path"],
                "workers_names": saved.get("workers_names"),
                "jobs_summary": saved.get("jobs_summary"),
                "jobs_count": saved.get("jobs_count"),
            }
        )
        results["saved"] = saved
        results["details"] = details
    except Exception as exc:
        results["steps"].append(
            {
                "step": "save_db",
                "ok": False,
                "error": str(exc),
                "trace": traceback.format_exc()[-800:],
            }
        )
        results["ok"] = False
        return results

    # 4) Checks
    details = results["details"]
    workers = details.get("workers") or []
    jobs = details.get("jobs") or []
    names_blob = " ".join(
        [
            str(details["report"].get("workers_names") or ""),
            " ".join(str(w.get("worker_name") or "") for w in workers),
            " ".join(str(w.get("members") or "") for w in workers),
        ]
    ).lower().replace("ё", "е")

    for hint in EXPECT_MEMBERS_HINTS:
        ok = hint in names_blob
        results["checks"].append(
            {"name": f"ФИО содержит «{hint}»", "ok": ok, "detail": names_blob[:200]}
        )
        results["ok"] = results["ok"] and ok

    brigade_ok = len(workers) == 1 and (workers[0].get("headcount") or 0) >= 2
    results["checks"].append(
        {
            "name": "Бригада = 1 worker, headcount >= 2",
            "ok": brigade_ok,
            "detail": f"workers={len(workers)}, headcount={workers[0].get('headcount') if workers else None}",
        }
    )
    results["ok"] = results["ok"] and brigade_ok

    jobs_ok = len(jobs) >= EXPECT_MIN_JOBS
    results["checks"].append(
        {
            "name": f"Работ в report_jobs >= {EXPECT_MIN_JOBS}",
            "ok": jobs_ok,
            "detail": f"jobs_count={len(jobs)}",
        }
    )
    results["ok"] = results["ok"] and jobs_ok

    vols = {j.get("volume") for j in jobs if j.get("volume") is not None}
    # key volumes that must appear
    key_vols = {7, 15, 25}
    hit = key_vols & vols
    vol_ok = key_vols.issubset(vols)
    results["checks"].append(
        {
            "name": "Ключевые объёмы 7 / 15 / 25 присутствуют",
            "ok": vol_ok,
            "detail": f"volumes={sorted(vols)}, hit={sorted(hit)}",
        }
    )
    results["ok"] = results["ok"] and vol_ok

    readable_ok = all(
        str(j.get("task_name") or "").strip() and str(j.get("worker_name") or "").strip()
        for j in jobs
    )
    results["checks"].append(
        {
            "name": "Все строки report_jobs имеют worker_name и task_name",
            "ok": readable_ok and bool(jobs),
            "detail": f"rows={len(jobs)}",
        }
    )
    results["ok"] = results["ok"] and readable_ok and bool(jobs)

    summary = str(details["report"].get("jobs_summary") or "")
    summary_ok = len(summary) > 20 and ("|" in summary or "—" in summary)
    results["checks"].append(
        {
            "name": "jobs_summary заполнен читаемым текстом",
            "ok": summary_ok,
            "detail": summary[:300],
        }
    )
    results["ok"] = results["ok"] and summary_ok

    return results


def write_docx(results: dict) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Отчёт по тестированию")
    set_run_font(r, size=20, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Модуль сохранения отчётов в БД (SQLite)\nМонтажПро")
    set_run_font(r, size=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Дата: {date.today().strftime('%d.%m.%Y')}  ·  "
        f"Итог: {'ПРОЙДЕНО' if results.get('ok') else 'ЕСТЬ ОШИБКИ'}"
    )
    set_run_font(r, size=11, bold=True)
    r.font.color.rgb = RGBColor(31, 157, 106) if results.get("ok") else RGBColor(214, 69, 69)

    add_heading(doc, "1. Цель", 1)
    add_para(
        doc,
        "Проверить цепочку: текст отчёта → разбор ИИ → расчёт по прайсу → сохранение в SQLite "
        "с читаемыми ФИО и наименованиями работ (таблицы reports, report_workers, report_jobs).",
    )

    add_heading(doc, "2. Входной текст", 1)
    p = doc.add_paragraph()
    run = p.add_run(SAMPLE_TEXT.strip())
    set_run_font(run, size=10)
    p.paragraph_format.space_after = Pt(8)

    add_heading(doc, "3. Шаги пайплайна", 1)
    step_rows = []
    for s in results.get("steps") or []:
        step_rows.append(
            [
                s.get("step"),
                "OK" if s.get("ok") else "FAIL",
                s.get("elapsed_sec", "—"),
                s.get("error")
                or s.get("workers_names")
                or s.get("grand_total")
                or s.get("id")
                or "",
            ]
        )
    t = doc.add_table(rows=1 + len(step_rows), cols=4)
    t.style = "Table Grid"
    fill_table(t, ["Шаг", "Статус", "Сек", "Детали"], step_rows)
    doc.add_paragraph()

    add_heading(doc, "4. Проверки читаемости БД", 1)
    check_rows = []
    for c in results.get("checks") or []:
        check_rows.append(
            [
                c.get("name"),
                "OK" if c.get("ok") else "FAIL",
                (c.get("detail") or "")[:180],
            ]
        )
    t2 = doc.add_table(rows=1 + len(check_rows), cols=3)
    t2.style = "Table Grid"
    fill_table(t2, ["Проверка", "Результат", "Детали"], check_rows)
    doc.add_paragraph()

    details = results.get("details") or {}
    workers = details.get("workers") or []
    jobs = details.get("jobs") or []

    add_heading(doc, "5. Данные в report_workers", 1)
    if workers:
        wrows = [
            [
                w.get("worker_no"),
                w.get("worker_name"),
                w.get("members"),
                w.get("headcount"),
                w.get("total"),
                w.get("per_person"),
            ]
            for w in workers
        ]
        tw = doc.add_table(rows=1 + len(wrows), cols=6)
        tw.style = "Table Grid"
        fill_table(
            tw,
            ["№", "ФИО / бригада", "members", "N", "Итого", "На чел."],
            wrows,
        )
    else:
        add_para(doc, "Нет строк workers.")
    doc.add_paragraph()

    add_heading(doc, "6. Данные в report_jobs", 1)
    if jobs:
        jrows = [
            [
                j.get("worker_name"),
                j.get("task_name"),
                j.get("volume"),
                j.get("unit"),
                j.get("unit_price") if j.get("unit_price") is not None else "—",
                j.get("amount") if j.get("amount") is not None else "—",
                "да" if j.get("matched") else "нет",
            ]
            for j in jobs
        ]
        tj = doc.add_table(rows=1 + len(jrows), cols=7)
        tj.style = "Table Grid"
        fill_table(
            tj,
            ["Сотрудник", "Работа", "Объём", "Ед.", "Тариф", "Сумма", "В прайсе"],
            jrows,
        )
    else:
        add_para(doc, "Нет строк jobs.")
    doc.add_paragraph()

    add_heading(doc, "7. Вывод", 1)
    if results.get("ok"):
        add_para(
            doc,
            "Тест модуля сохранения в БД пройден: бригада распознана, работы записаны "
            "в report_jobs с читаемыми именами и объёмами, сводка jobs_summary заполнена.",
            bold=True,
        )
    else:
        add_para(
            doc,
            "Тест завершён с ошибками — см. таблицы шагов и проверок. "
            "При сбое разбора ИИ проверьте Ollama и таймаут; при сбое матчинга — прайс/алиасы.",
            bold=True,
        )

    saved = results.get("saved") or {}
    add_para(
        doc,
        f"Тестовая БД: {saved.get('db_path') or TEST_DB}. "
        f"Файл отчёта: {OUT_DOCX.name}",
        size=10,
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


def main():
    print("Running DB save pipeline test...")
    results = run_pipeline()
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    # trim huge blobs for json dump
    dump = {
        k: v
        for k, v in results.items()
        if k not in ("report", "calculation")
    }
    dump["report_preview"] = {
        "workers": [
            {
                "raw": w.get("raw_worker_name"),
                "members": w.get("members"),
                "jobs": len(w.get("performed_jobs") or []),
            }
            for w in (results.get("report") or {}).get("workers") or []
        ]
    }
    OUT_JSON.write_text(json.dumps(dump, ensure_ascii=False, indent=2), encoding="utf-8")
    path = write_docx(results)
    print("ok=", results.get("ok"))
    for c in results.get("checks") or []:
        print(("PASS" if c["ok"] else "FAIL"), c["name"], "|", c.get("detail", "")[:100])
    print("docx", path)
    print("json", OUT_JSON)
    print("db", TEST_DB)


if __name__ == "__main__":
    main()
